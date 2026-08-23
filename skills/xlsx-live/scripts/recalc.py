#!/usr/bin/env python3
"""実Excel(xlwings)でxlsxの数式を再計算し、エラーをJSONで報告する。

openpyxlで作成・編集したファイルは数式が文字列として入るだけで
計算値を持たない。実Excelに開かせて再計算・保存し、その後
openpyxlのdata_onlyで全セルを走査してExcelエラー文字列を数える。

初回実行時はmacOSが「ターミナルにExcelの制御を許可しますか」という
オートメーション許可ダイアログを出すことがあり、応答するまで処理が
止まる。タイムアウトした場合はダイアログの有無を確認すること。

usage: python recalc.py <file.xlsx> [timeout_seconds]  (default 120)
"""

from __future__ import annotations

import json
import signal
import sys
from pathlib import Path

ERROR_STRINGS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")

TIMEOUT_HINT = (
    "タイムアウトしました。初回はオートメーション許可ダイアログや"
    "Excelの初回起動画面が応答待ちになっている可能性があります。"
    "画面のダイアログに応答するか、システム設定 > プライバシーとセキュリティ >"
    " オートメーション で端末アプリにMicrosoft Excelの制御を許可してください。"
)


class RecalcTimeout(Exception):
    pass


def recalc_with_excel(path: Path, timeout: int) -> None:
    import xlwings as xw

    def on_timeout(signum, frame):
        raise RecalcTimeout(TIMEOUT_HINT)

    signal.signal(signal.SIGALRM, on_timeout)
    signal.alarm(timeout)
    app = None
    try:
        app = xw.App(visible=False)
        wb = app.books.open(str(path))
        app.api.calculate()  # 全ブック再計算
        wb.save()
        wb.close()
    finally:
        signal.alarm(0)
        if app is not None:
            try:
                app.quit()
            except Exception:
                pass


def scan_errors(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    formula_wb = load_workbook(path)

    total_formulas = 0
    summary: dict[str, dict] = {}
    for ws in formula_wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    total_formulas += 1
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v in ERROR_STRINGS:
                    entry = summary.setdefault(v, {"count": 0, "locations": []})
                    entry["count"] += 1
                    if len(entry["locations"]) < 20:
                        entry["locations"].append(f"{ws.title}!{cell.coordinate}")

    total_errors = sum(e["count"] for e in summary.values())
    result = {
        "status": "errors_found" if total_errors else "success",
        "total_errors": total_errors,
        "total_formulas": total_formulas,
    }
    if summary:
        result["error_summary"] = summary
    return result


def main() -> int:
    if len(sys.argv) < 2:
        print(json.dumps({"status": "usage_error", "message": "usage: recalc.py <file.xlsx>"}))
        return 2
    path = Path(sys.argv[1]).resolve()
    if not path.exists():
        print(json.dumps({"status": "not_found", "path": str(path)}))
        return 2
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 120
    try:
        recalc_with_excel(path, timeout)
    except RecalcTimeout as exc:
        print(json.dumps({"status": "timeout", "error": str(exc)}, ensure_ascii=False))
        return 1
    except Exception as exc:  # Excel不在・xlwings未導入など
        print(json.dumps({"status": "recalc_failed", "error": str(exc)}, ensure_ascii=False))
        return 1
    result = scan_errors(path)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "success" else 1


if __name__ == "__main__":
    sys.exit(main())
