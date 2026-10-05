#!/usr/bin/env python3
"""統合版の生成 / 末尾へのページ追加。

初回生成（S01..Snn を順に結合。1枚目のpptxを土台にする）:
  merge_deck.py build --slides-dir パワポ --out 統合版/提案書_統合版.pptx

末尾に1枚追加（★autofitを壊さない安全な方法。既存ページには触らない）:
  merge_deck.py append --deck 統合版/提案書_統合版.pptx \
                       --slide S18_ランドスケープ/S18_ランドスケープ.pptx

⚠️ build は全ページを組み直すため、プレースホルダーの継承が切れて
   タイトル/メッセージラインのフォントサイズが飛ぶことがある。
   既存デッキがあるなら build ではなく append / patch_text.py を使う。
   build を使ったら必ず visual_qa.py で全ページを確認する。
   詳細は references/pitfalls.md の「1. 全ページ再結合は autofit を壊す」。
"""
import argparse, copy, glob, io, os, shutil, sys
from datetime import datetime
from pptx import Presentation
from pptx.oxml.ns import qn


def backup(path):
    d = os.path.join(os.path.dirname(os.path.abspath(path)), "old")
    os.makedirs(d, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    base, ext = os.path.splitext(os.path.basename(path))
    dst = os.path.join(d, f"{base}_{ts}{ext}")
    n = 2
    while os.path.exists(dst):          # 同一分内の衝突で退避を失わない
        dst = os.path.join(d, f"{base}_{ts}_{n}{ext}")
        n += 1
    shutil.copy(path, dst)
    print(f"退避: {dst}")


def match_layout(dst_pres, src_slide, fallback):
    """元スライドと同名のレイアウトを結合先から探す。無ければ fallback。
    表紙のレイアウトを全ページに使うと背景図が乗り、プレースホルダーの継承も切れる。"""
    name = src_slide.slide_layout.name
    for lay in dst_pres.slide_layouts:
        if lay.name == name:
            return lay
    return fallback


_R_ATTRS = (qn("r:embed"), qn("r:link"), qn("r:id"))


def _remap(el, src_slide, new, rmap):
    """図形 XML 内の r:embed / r:link / r:id を、結合先スライドのリレーションに張り替える。
    張り替えないと画像・リンクの参照先が無くなり、PowerPoint で画像が消える（2026-09-19 修正）。"""
    for node in el.iter():
        for at in _R_ATTRS:
            rid = node.get(at)
            if not rid:
                continue
            if rid not in rmap:
                rel = src_slide.part.rels[rid]
                if rel.is_external:
                    rmap[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                elif rel.reltype.endswith("/image"):
                    # 別パッケージの画像パーツは名前が衝突するので、blob から取り込み直す
                    _, rmap[rid] = new.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                else:
                    sys.exit(f"未対応のリレーション: {rel.reltype}（{src_slide.part.partname}）")
            node.set(at, rmap[rid])


def copy_shapes(dst_pres, src_slide, layout):
    layout = match_layout(dst_pres, src_slide, layout)
    new = dst_pres.slides.add_slide(layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    rmap = {}
    for sh in src_slide.shapes:
        el = copy.deepcopy(sh._element)
        _remap(el, src_slide, new, rmap)
        new.shapes._spTree.append(el)
    bg = src_slide._element.cSld.find(qn("p:bg"))   # スライド固有の背景も複製
    if bg is not None:
        el = copy.deepcopy(bg)
        _remap(el, src_slide, new, rmap)
        new._element.cSld.insert(0, el)
    return new


def cmd_build(a):
    dirs = sorted(d for d in glob.glob(os.path.join(a.slides_dir, "S*")) if os.path.isdir(d))
    files = []
    for d in dirs:
        hits = [f for f in glob.glob(os.path.join(d, "*.pptx"))
                if os.sep + "old" + os.sep not in f and not os.path.basename(f).startswith("~$")]
        if hits:
            files.append(sorted(hits)[0])
    if not files:
        sys.exit("結合対象の pptx が見つからない")
    print(f"結合対象 {len(files)} 枚:")
    for f in files:
        print("  ", f)

    if os.path.exists(a.out):
        backup(a.out)
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)

    base = Presentation(files[0])
    lst = base.slides._sldIdLst
    for sid in list(lst)[1:]:
        lst.remove(sid)
    for f in files[1:]:
        src = Presentation(f)
        for sl in src.slides:
            copy_shapes(base, sl, base.slides[-1].slide_layout)
    base.save(a.out)
    print(f"\n生成: {a.out}（{len(Presentation(a.out).slides)}ページ）")
    print("⚠️ visual_qa.py で全ページのレンダリングを確認すること")


def cmd_append(a):
    backup(a.deck)
    dst = Presentation(a.deck)
    src = Presentation(a.slide)
    before = len(dst.slides)
    layout = dst.slides[-1].slide_layout   # 既存最終ページのレイアウトを継承
    for sl in src.slides:
        copy_shapes(dst, sl, layout)
    dst.save(a.deck)
    print(f"{before} → {len(Presentation(a.deck).slides)} ページ")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build"); b.add_argument("--slides-dir", required=True); b.add_argument("--out", required=True)
    p = sub.add_parser("append"); p.add_argument("--deck", required=True); p.add_argument("--slide", required=True)
    a = ap.parse_args()
    {"build": cmd_build, "append": cmd_append}[a.cmd](a)


if __name__ == "__main__":
    main()
