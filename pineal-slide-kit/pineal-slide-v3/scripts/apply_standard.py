#!/usr/bin/env python3
"""既存デッキの「中面A」ページに pineal 標準を後から適用する。

標準（2026-09 確定 / references/design-system.md）:
  - タイトル: top 0.346 / height 0.471（28pt 前提）
  - リード文: 行頭に ✓（Wingdings の "ü"）／ marL 285750 / indent -285750

  apply_standard.py --file deck.pptx
  apply_standard.py --file deck.pptx --dry-run          # 対象の確認のみ
  apply_standard.py --file deck.pptx --layout 中面B     # 別レイアウトを対象にする
  apply_standard.py --file deck.pptx --no-bullet        # タイトル位置だけ揃える
  apply_standard.py --file deck.pptx --skip 2           # 手作業済みのページを除く

新規スライドはテンプレ側に反映済みなので、このスクリプトは不要
（既存デッキの追いつき用）。タイトルのpt自体は set_font.py title-size で変える。
"""
import argparse, os, shutil, sys
from datetime import datetime
from lxml import etree
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.ns import qn

NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
PPR = (f'<a:pPr {NS} marL="285750" indent="-285750">'
       '<a:buFont typeface="Wingdings" pitchFamily="2" charset="2"/>'
       '<a:buChar char="ü"/>'
       '</a:pPr>')
# 4値すべてを設定する。top/height だけ指定すると <a:xfrm> が新規作成され、
# left/width の「レイアウトからの継承」が切れて 0 になる（タイトルが縦1列に潰れる）。
TITLE_L, TITLE_T, TITLE_W, TITLE_H = 0.602, 0.346, 12.128, 0.471


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--layout", default="中面A")
    ap.add_argument("--skip", type=int, nargs="*", default=[], help="除外するページ番号")
    ap.add_argument("--no-bullet", action="store_true")
    ap.add_argument("--no-title", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    p = Presentation(a.file)
    hits = []
    for i, s in enumerate(p.slides, 1):
        if s.slide_layout.name != a.layout or i in a.skip:
            continue
        title = msg = None
        for sh in s.shapes:
            if sh.is_placeholder:
                if sh.placeholder_format.idx == 0: title = sh
                elif sh.placeholder_format.idx == 1: msg = sh
        acts = []
        if title is not None and not a.no_title:
            if not a.dry_run:
                title.left   = Inches(TITLE_L)
                title.top    = Inches(TITLE_T)
                title.width  = Inches(TITLE_W)
                title.height = Inches(TITLE_H)
            acts.append("タイトル位置/サイズ")
        if msg is not None and not a.no_bullet:
            if not a.dry_run:
                for para in msg.text_frame.paragraphs:
                    el = para._p
                    old = el.find(qn('a:pPr'))
                    if old is not None:
                        el.remove(old)
                    el.insert(0, etree.fromstring(PPR))
            acts.append("リード文に✓")
        hits.append((i, acts))

    if not hits:
        sys.exit(f"レイアウト {a.layout!r} のページが無い")
    for i, acts in hits:
        print(f"  p{i}: {' + '.join(acts)}")

    if a.dry_run:
        print("\n--dry-run: 保存していない")
        return

    d = os.path.join(os.path.dirname(os.path.abspath(a.file)), "old")
    os.makedirs(d, exist_ok=True)
    base, ext = os.path.splitext(os.path.basename(a.file))
    bk = os.path.join(d, f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}_beforestd{ext}")
    shutil.copy(a.file, bk)
    print(f"\n退避: {bk}")
    p.save(a.file)
    print(f"保存: {a.file}（{len(hits)}ページ）")


if __name__ == "__main__":
    main()
