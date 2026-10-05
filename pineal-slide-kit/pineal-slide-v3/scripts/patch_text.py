#!/usr/bin/env python3
"""統合版のテキストを文字列一致で直接置換する。再結合しないので autofit を壊さない。

  patch_text.py --deck 統合版/提案書_統合版.pptx \
      --replace "旧テキスト" "新テキスト" \
      --replace "旧2" "新2"

各置換が「ちょうど1件ヒット」することを検証し、0件または複数件なら中断する。
--allow-multi で複数件を許可。--dry-run で確認のみ。
run 単位で一致させるため、run が分割されている場合は一致しない
（そのときは対象shapeを特定して run のテキストを直接指定する）。
"""
import argparse, os, shutil, sys
from datetime import datetime
from pptx import Presentation


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", required=True)
    ap.add_argument("--replace", nargs=2, action="append", required=True, metavar=("OLD", "NEW"))
    ap.add_argument("--allow-multi", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    p = Presentation(a.deck)
    hits = {old: [] for old, _ in a.replace}
    table = dict(a.replace)

    for i, s in enumerate(p.slides, 1):
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.text in table:
                        hits[r.text].append(i)
                        if not a.dry_run:
                            r.text = table[r.text]

    ng = False
    for old, _ in a.replace:
        n = len(hits[old])
        mark = "OK " if (n == 1 or (a.allow_multi and n > 0)) else "NG "
        if mark == "NG ":
            ng = True
        print(f"{mark}{n}件 p{hits[old]} | {old[:50]}")
    if ng:
        sys.exit("\n0件または想定外の複数件があるため中断（--allow-multi で許可可）")

    if a.dry_run:
        print("\n--dry-run: 保存していない")
        return

    d = os.path.join(os.path.dirname(os.path.abspath(a.deck)), "old")
    os.makedirs(d, exist_ok=True)
    base, ext = os.path.splitext(os.path.basename(a.deck))
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    bk = os.path.join(d, f"{base}_{ts}{ext}")
    n = 2
    while os.path.exists(bk):
        bk = os.path.join(d, f"{base}_{ts}_{n}{ext}")
        n += 1
    shutil.copy(a.deck, bk)
    print(f"\n退避: {bk}")
    p.save(a.deck)
    print(f"保存: {a.deck}")


if __name__ == "__main__":
    main()
