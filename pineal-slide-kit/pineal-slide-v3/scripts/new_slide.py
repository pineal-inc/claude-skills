#!/usr/bin/env python3
"""テンプレから1枚取り出してスライド単位のフォルダ＋pptxを起こす。

使い方:
  new_slide.py --out "S03_プロジェクト理解/S03_プロジェクト理解.pptx"
  new_slide.py --out "..." --source-slide 10        # As-Is/To-Be の型を流用
  new_slide.py --out "..." --keep-body              # ボディ図形も残す（型を流用するとき）

既定では タイトル / メッセージライン / ページ番号 だけを残し、ボディは削除する。
--source-slide は references/template-catalog.md の番号（1始まり）。
"""
import argparse, os, shutil, sys
from pptx import Presentation

DEFAULT_TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "assets", "pineal_temp.pptx")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True, help="出力先 pptx のパス")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    ap.add_argument("--source-slide", type=int, default=8, help="流用するテンプレのスライド番号(1始まり)")
    ap.add_argument("--keep-body", action="store_true", help="ボディ図形を残す")
    ap.add_argument("--title", default=None)
    ap.add_argument("--message", default=None)
    a = ap.parse_args()

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    os.makedirs(os.path.join(os.path.dirname(os.path.abspath(a.out)), "old"), exist_ok=True)

    shutil.copy(a.template, a.out)
    p = Presentation(a.out)

    idx = a.source_slide - 1
    if not (0 <= idx < len(p.slides)):
        sys.exit(f"source-slide が範囲外: 1〜{len(p.slides)}")

    # 目的のスライド以外を削除
    lst = p.slides._sldIdLst
    for i, sid in enumerate(list(lst)):
        if i != idx:
            rid = sid.rId
            lst.remove(sid)
            p.part.drop_rel(rid)   # 孤児パーツを残すと再保存時に slide1.xml が重複して壊れる

    s = p.slides[0]
    if not a.keep_body:
        # プレースホルダー（タイトル/メッセージ/ページ番号）以外を削除
        for sh in list(s.shapes):
            if not sh.is_placeholder:
                sh._element.getparent().remove(sh._element)

    def set_ph(order, text):
        phs = [sh for sh in s.shapes if sh.is_placeholder and sh.has_text_frame]
        phs.sort(key=lambda x: (x.top or 0))
        if order < len(phs) and text is not None:
            tf = phs[order].text_frame
            if tf.paragraphs and tf.paragraphs[0].runs:
                tf.paragraphs[0].runs[0].text = text
            else:
                tf.text = text

    set_ph(0, a.title)
    set_ph(1, a.message)

    p.save(a.out)
    print(f"作成: {a.out}")
    print(f"  流用元: テンプレ スライド{a.source_slide} / ボディ: {'保持' if a.keep_body else '削除'}")
    for sh in Presentation(a.out).slides[0].shapes:
        t = sh.text_frame.text.replace("\n", " / ")[:44] if sh.has_text_frame else ""
        print(f"  id={sh.shape_id} ph={sh.is_placeholder} | {t}")


if __name__ == "__main__":
    main()
